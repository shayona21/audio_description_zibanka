"""Read the standardized workbook by header name, independently of V1."""

import re
from datetime import time
from pathlib import Path

from openpyxl import load_workbook

from .models import DialogueRow

REQUIRED_HEADERS = (
    "EP NO", "TCR IN", "TCR OUT", "CHARACTERS", "VOICE ID", "TARGET DIALOGUE",
)
TEMPLATE_HEADERS = (
    "EP NO", "SR NO", "TCR IN", "TCR OUT", "CHARACTERS", "VOICE ID",
    "TARGET DIALOGUE", "DUBBING STATUS",
)
SUPPORTED_FPS = (24, 25, 30, 50, 60)


class WorkbookValidationError(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("\n".join(errors))


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def timecode_to_ms(value, fps=25):
    if fps not in SUPPORTED_FPS:
        raise ValueError(f"Frame rate must be one of {SUPPORTED_FPS}")
    if isinstance(value, time):
        return round(((value.hour * 60 + value.minute) * 60 + value.second) * 1000
                     + value.microsecond / 1000)
    text = cell_text(value)
    if not re.fullmatch(r"\d{2,}:\d{2}:\d{2}:\d{2}", text):
        raise ValueError(f"{text!r}: expected HH:MM:SS:FF (for example 00:00:18:08)")
    hours, minutes, seconds, frames = map(int, text.split(":"))
    if minutes >= 60 or seconds >= 60 or frames >= fps:
        raise ValueError(f"Invalid timecode {text!r} at {fps} fps")
    return round(((hours * 60 + minutes) * 60 + seconds + frames / fps) * 1000)


def parse_excel(path, fps=25, sheet_name=None):
    if fps not in SUPPORTED_FPS:
        raise WorkbookValidationError([f"Frame rate must be one of {SUPPORTED_FPS}"])
    if Path(path).suffix.lower() != ".xlsx":
        raise WorkbookValidationError(["V2 requires an .xlsx workbook"])
    # Read formulas as formulas so they cannot silently become empty cached values.
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise WorkbookValidationError([f"Worksheet {sheet_name!r} was not found"])
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        source = sheet.iter_rows(values_only=True)
        headers = [" ".join(cell_text(v).upper().split()) for v in next(source, ())]
        errors = []
        for name in REQUIRED_HEADERS:
            if headers.count(name) == 0:
                errors.append(f"Missing required header: {name}")
            elif headers.count(name) > 1:
                errors.append(f"Duplicate header: {name}")
        if errors:
            raise WorkbookValidationError(errors)
        columns = {name: headers.index(name) for name in REQUIRED_HEADERS}
        rows, voices = [], {}
        for number, values in enumerate(source, start=2):
            cells = {name: values[index] if index < len(values) else None
                     for name, index in columns.items()}
            if not any(cell_text(value) for value in cells.values()):
                continue
            # A missing voice marks a row that should not be synthesized.
            if not cell_text(cells["VOICE ID"]):
                continue
            missing = [name for name, value in cells.items() if not cell_text(value)]
            if missing:
                errors.append(f"Row {number}: missing {', '.join(missing)}")
                continue
            formulas = [name for name, value in cells.items()
                        if isinstance(value, str) and value.startswith("=")]
            if formulas:
                errors.append(f"Row {number}: paste values instead of formulas in {', '.join(formulas)}")
                continue
            try:
                start = timecode_to_ms(cells["TCR IN"], fps)
                end = timecode_to_ms(cells["TCR OUT"], fps)
                if end <= start:
                    raise ValueError("TCR OUT must be later than TCR IN")
            except ValueError as exc:
                errors.append(f"Row {number}: {exc}")
                continue
            row = DialogueRow(number, cell_text(cells["EP NO"]),
                              cell_text(cells["CHARACTERS"]), start, end,
                              cell_text(cells["TARGET DIALOGUE"]),
                              cell_text(cells["VOICE ID"]))
            if row.group in voices and voices[row.group] != row.voice_id:
                errors.append(f"Row {number}: conflicting VOICE ID for episode {row.episode}, character {row.character}")
            voices[row.group] = row.voice_id
            rows.append(row)
        if errors:
            raise WorkbookValidationError(errors)
        if not rows:
            raise WorkbookValidationError(["No dialogue rows found"])
        return rows
    finally:
        workbook.close()
