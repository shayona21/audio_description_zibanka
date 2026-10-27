import csv
import io
import json
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import Mock, patch
from urllib.error import HTTPError
import wave
import zipfile

from openpyxl import Workbook, load_workbook
from pydub import AudioSegment
from pydub.generators import Sine

from v2.app import create_app
from v2.audio_speed import (FIT_TOLERANCE_MS, MAX_FIT_SPEED, MIN_FIT_SPEED, change_speed, fit_speed, needed_speed,
                            retime, speech_span)
from v2.excel_parser import TEMPLATE_HEADERS, WorkbookValidationError, parse_excel, timecode_to_ms
from v2.exporter import format_tcr, track_names, timeline_names
from v2.models import DialogueRow
from v2.pipeline import run_pipeline
from v2.timing import place_dialogue
from v2.tts_client import MODEL_ID, ElevenLabsClient, prepare_speech_text


def wav_bytes(duration_ms=200, value=1000):
    output = io.BytesIO()
    with wave.open(output, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(24000)
        wav.writeframes(struct.pack("<h", value) * (duration_ms * 24))
    return output.getvalue()


def dialogue(episode="001", character="Alice", start="00:00:01:00", end="00:00:02:00",
             text="Hello", voice="alice-voice"):
    return [episode, 1, start, end, character, voice, text, "Done"]


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def workbook(self, rows, headers=TEMPLATE_HEADERS):
        path = self.root / "input.xlsx"
        book = Workbook()
        book.active.append(list(headers))
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def test_headers_are_mapped_by_name_and_ignored_columns_do_not_filter_rows(self):
        headers = list(reversed(TEMPLATE_HEADERS))
        headers[headers.index("EP NO")] = " ep  no "
        path = self.workbook([list(reversed(dialogue()))], headers)
        rows = parse_excel(path)
        self.assertEqual(rows[0].episode, "001")
        self.assertEqual(rows[0].character, "Alice")
        self.assertEqual(rows[0].voice_id, "alice-voice")
        self.assertEqual(rows[0].text, "Hello")
        self.assertEqual(rows[0].row_number, 2)

    def test_validation_reports_all_bad_rows_before_any_api_call(self):
        path = self.workbook([
            dialogue(),
            dialogue(character=""),
            dialogue(start="00:00:01:25"),
            dialogue(end="00:00:00:00"),
        ])
        synthesize = Mock()
        with self.assertRaises(WorkbookValidationError) as caught:
            run_pipeline(path, self.root / "out.zip", synthesize=synthesize)
        self.assertIn("Row 3", str(caught.exception))
        self.assertIn("Row 4", str(caught.exception))
        self.assertIn("Row 5", str(caught.exception))
        synthesize.assert_not_called()

    def test_conflicting_voices_are_rejected_with_source_row(self):
        path = self.workbook([dialogue(), dialogue(voice="different")])
        with self.assertRaisesRegex(WorkbookValidationError, "Row 3: conflicting VOICE ID"):
            parse_excel(path)

    def test_missing_duplicate_headers_empty_sheet_and_formulas(self):
        for headers, rows, message in [
            (TEMPLATE_HEADERS[:-2], [], "Missing required header: TARGET DIALOGUE"),
            (TEMPLATE_HEADERS + ("EP NO",), [], "Duplicate header"),
            (TEMPLATE_HEADERS, [], "No dialogue rows"),
            (TEMPLATE_HEADERS, [dialogue(text='="Hello"')], "paste values"),
        ]:
            with self.subTest(message=message):
                path = self.workbook(rows, headers)
                with self.assertRaisesRegex(WorkbookValidationError, message):
                    parse_excel(path)

    def test_blank_rows_preserve_excel_row_numbers(self):
        path = self.workbook([[None] * 8, dialogue()])
        self.assertEqual(parse_excel(path)[0].row_number, 3)

    def test_frame_conversion_and_sheet_selection(self):
        self.assertEqual(timecode_to_ms("00:00:18:08"), 18320)
        self.assertEqual(timecode_to_ms("00:00:00:12", 24), 500)
        path = self.workbook([dialogue()])
        with self.assertRaisesRegex(WorkbookValidationError, "was not found"):
            parse_excel(path, sheet_name="Missing")
        self.assertEqual(len(parse_excel(path, sheet_name="Sheet")), 1)

    def test_output_grouping_voice_routing_alignment_and_overrun_preservation(self):
        # Deliberately unsorted; different characters may start simultaneously.
        path = self.workbook([
            dialogue(start="00:00:03:00", end="00:00:03:04", text="late"),
            dialogue(character="Bob", start="00:00:01:00", text="bob", voice="bob-voice"),
            dialogue(text="early"),
            dialogue(episode="002", start="00:00:00:00", end="00:00:01:00", text="episode2"),
        ])
        calls = []

        def synthesize(text, voice):
            calls.append((text, voice))
            return wav_bytes(400 if text == "late" else 200)

        output = self.root / "result.zip"
        summary = run_pipeline(path, output, synthesize=synthesize, progress=lambda _: None,
                               fit_lines=False)
        self.assertEqual(calls, [("late", "alice-voice"), ("bob", "bob-voice"),
                                 ("early", "alice-voice"), ("episode2", "alice-voice")])
        self.assertEqual(summary["tracks"], 3)
        self.assertEqual(summary["flagged_rows"], 1)
        with zipfile.ZipFile(output) as archive:
            self.assertEqual(set(archive.namelist()), {
                "tracks/Episode_001_Alice.wav", "tracks/Episode_001_Bob.wav", "tracks/Episode_002_Alice.wav",
                "lines/001_Alice_2.wav", "lines/001_Bob_3.wav", "lines/001_Alice_4.wav", "lines/002_Alice_5.wav",
                "Episode_001_timeline.csv", "Episode_002_timeline.csv", "processing_report.csv", "script.xlsx"})
            self.assertEqual(archive.read("script.xlsx"), path.read_bytes())
            for name in ("Episode_001_Alice.wav", "Episode_001_Bob.wav"):
                with wave.open(io.BytesIO(archive.read(f"tracks/{name}"))) as wav:
                    self.assertEqual(wav.getnframes(), 3400 * 24)
                    self.assertEqual(wav.getframerate(), 24000)
                    self.assertEqual(wav.getnchannels(), 1)
                    pcm = wav.readframes(wav.getnframes())
                    first_start = 1000
                    self.assertEqual(pcm[:first_start * 24 * 2], bytes(first_start * 24 * 2))
                    self.assertEqual(struct.unpack_from("<h", pcm, first_start * 24 * 2)[0], 1000)
                    if "Alice" in name:
                        self.assertEqual(struct.unpack_from("<h", pcm, 3300 * 24 * 2)[0], 1000)
            with wave.open(io.BytesIO(archive.read("tracks/Episode_002_Alice.wav"))) as wav:
                self.assertEqual(wav.getnframes(), 1000 * 24)
            report = list(csv.DictReader(io.StringIO(archive.read("processing_report.csv").decode("utf-8-sig"))))
            self.assertEqual(report[0]["overrun_amnt_ms"], "240")
            self.assertEqual(report[1]["status"], "OK")

    def test_same_character_collision_uses_actual_end_not_script_end(self):
        path = self.workbook([dialogue(), dialogue(start="00:00:01:04")])
        output = self.root / "overlap.zip"
        run_pipeline(path, output, synthesize=lambda *_: wav_bytes(400), progress=lambda _: None,
                     fit_lines=False)
        with zipfile.ZipFile(output) as archive:
            rows = list(csv.DictReader(io.StringIO(archive.read("processing_report.csv").decode("utf-8-sig"))))
        self.assertEqual(rows[1]["status"], "OK")
        self.assertEqual(rows[1]["adjusted_start_TCR"], "00:00:01:10")
        self.assertEqual(rows[1]["start_TCR"], "00:00:01:04")
        self.assertEqual(rows[1]["overrun_amnt_ms"], "0")

    def test_character_local_delays_match_audio_and_combined_script_exports(self):
        # Unsorted characters and episodes still have independent schedules.
        path = self.workbook([
            dialogue(character="A", start="00:00:40:00", end="00:00:42:00", text="Why?", voice="a"),
            dialogue(episode="002", character="A", start="00:00:01:00", end="00:00:02:00",
                     text="Other episode", voice="a"),
            dialogue(character="B", start="00:00:11:00", end="00:00:14:00", text="[laughs]", voice="b"),
            dialogue(character="A", start="00:00:00:00", end="00:00:10:00", text="Walk?", voice="a"),
        ])
        durations = {"Walk?": 45000, "[laughs]": 3000, "Why?": 2000, "Other episode": 1000}
        values = {"Walk?": 1000, "[laughs]": 2000, "Why?": 3000, "Other episode": 4000}
        synthesize = Mock(side_effect=lambda text, voice: wav_bytes(durations[text], values[text]))
        output = self.root / "character-timing.zip"
        run_pipeline(path, output, synthesize=synthesize, progress=lambda _: None,
                     fit_lines=False)
        self.assertEqual(synthesize.call_count, 4)
        expected = {
            "Episode_001_A": [(0, 45000, "Walk?", 0), (45000, 47000, "Why?", 5000)],
            "Episode_001_B": [(11000, 14000, "[laughs]", 0)],
            "Episode_002_A": [(1000, 2000, "Other episode", 0)],
        }
        with zipfile.ZipFile(output) as archive:
            for stem, expected_rows in expected.items():
                with self.subTest(track=stem):
                    with wave.open(io.BytesIO(archive.read(f"tracks/{stem}.wav"))) as audio:
                        duration = 47000 if stem.startswith("Episode_001") else 2000
                        self.assertEqual(audio.getnframes(), duration * 24)
                        actual = audio.readframes(audio.getnframes())
                    expected_pcm = bytearray(duration * 24 * 2)
                    for start, end, text, _ in expected_rows:
                        expected_pcm[start * 24 * 2:end * 24 * 2] = struct.pack("<h", values[text]) * ((end - start) * 24)
                    self.assertEqual(actual, bytes(expected_pcm))
            timeline = list(csv.reader(io.StringIO(
                archive.read("Episode_001_timeline.csv").decode("utf-8-sig")
            )))
            self.assertEqual(timeline, [
                ["Dialogue", "Original time start", "Original time end",
                 "Adjusted time start", "Adjusted time end", "Delay"],
                ["A: Walk?", "00:00:00.000", "00:00:10.000", "00:00:00.000", "00:00:45.000", "0.000 s"],
                ["B: [laughs]", "00:00:11.000", "00:00:14.000", "00:00:11.000", "00:00:14.000", "0.000 s"],
                ["A: Why?", "00:00:40.000", "00:00:42.000", "00:00:45.000", "00:00:47.000", "5.000 s"],
            ])
            other = list(csv.DictReader(io.StringIO(
                archive.read("Episode_002_timeline.csv").decode("utf-8-sig")
            )))
            self.assertEqual(len(other), 1)
            self.assertEqual(other[0]["Dialogue"], "A: Other episode")
            self.assertEqual(other[0]["Delay"], "0.000 s")
            report = list(csv.DictReader(io.StringIO(
                archive.read("processing_report.csv").decode("utf-8-sig")
            )))
            self.assertEqual(report[0]["start_TCR"], "00:00:40:00")
            self.assertEqual(report[0]["adjusted_start_TCR"], "00:00:45:00")
            self.assertEqual(report[0]["duration"], "2000")
            self.assertEqual(report[0]["overrun_amnt_ms"], "0")

    def test_available_silence_absorbs_delays_within_each_character(self):
        path = self.workbook([
            dialogue(start="00:00:00:00", end="00:00:10:00", text="first"),
            dialogue(start="00:00:11:00", end="00:00:14:00", text="short"),
            dialogue(start="00:00:20:00", end="00:00:22:00", text="long"),
            dialogue(character="Bob", voice="b", start="00:00:12:00", end="00:00:13:00", text="bob"),
            dialogue(start="00:00:25:00", end="00:00:27:00", text="last"),
        ])
        durations = {"first": 15000, "short": 1000, "long": 6000, "bob": 500, "last": 500}
        output = self.root / "recovered-timing.zip"
        run_pipeline(path, output, synthesize=lambda text, voice: wav_bytes(durations[text]), progress=lambda _: None,
                     fit_lines=False)
        with zipfile.ZipFile(output) as archive:
            records = list(csv.DictReader(io.StringIO(
                archive.read("processing_report.csv").decode("utf-8-sig")
            )))
            self.assertEqual([r["adjusted_start_TCR"] for r in records],
                             ["00:00:00:00", "00:00:15:00", "00:00:20:00", "00:00:12:00", "00:00:26:00"])
            self.assertEqual([r["adjusted_end_TCR"] for r in records],
                             ["00:00:15:00", "00:00:16:00", "00:00:26:00", "00:00:12:13", "00:00:26:13"])
            self.assertEqual([int(r["overrun_amnt_ms"]) for r in records], [5000, 0, 4000, 0, 0])
            timeline = list(csv.DictReader(io.StringIO(archive.read("Episode_001_timeline.csv").decode("utf-8-sig"))))
            self.assertEqual([r["Dialogue"] for r in timeline],
                             ["Alice: first", "Bob: bob", "Alice: short", "Alice: long", "Alice: last"])
            self.assertEqual([r["Delay"] for r in timeline], ["0.000 s", "0.000 s", "4.000 s", "0.000 s", "1.000 s"])
            for name in ("Episode_001_Alice.wav", "Episode_001_Bob.wav"):
                with wave.open(io.BytesIO(archive.read(f"tracks/{name}"))) as audio:
                    self.assertEqual(audio.getnframes(), 27000 * 24)

    def test_timeline_preserves_millisecond_precision_and_csv_dialogue(self):
        text = '+Hello, "world"\nनमस्ते'
        path = self.workbook([
            dialogue(start="00:00:00:01", end="00:00:01:01", text=text),
            dialogue(start="00:00:01:12", end="00:00:02:12", text="second"),
        ])
        output = self.root / "precision.zip"
        run_pipeline(path, output, fps=24, synthesize=lambda *_: wav_bytes(1501), progress=lambda _: None,
                     fit_lines=False)
        with zipfile.ZipFile(output) as archive:
            records = list(csv.DictReader(io.StringIO(
                archive.read("Episode_001_timeline.csv").decode("utf-8-sig")
            )))
        self.assertEqual(records[0]["Adjusted time start"], "00:00:00.042")
        self.assertEqual(records[0]["Adjusted time end"], "00:00:01.543")
        self.assertEqual(records[1]["Adjusted time start"], "00:00:01.543")
        self.assertEqual(records[1]["Delay"], "0.043 s")
        self.assertEqual(records[0]["Dialogue"], "Alice: " + text)

    def test_only_actual_collisions_delay_a_character(self):
        rows = [
            DialogueRow(2, "1", "A", 0, 10000, "first", "shared-voice"),
            DialogueRow(3, "1", "A", 40000, 42000, "next", "shared-voice"),
            DialogueRow(4, "1", "A", 44000, 45000, "last", "shared-voice"),
            DialogueRow(5, "1", "B", 11000, 14000, "other character", "shared-voice"),
            DialogueRow(6, "2", "A", 1000, 2000, "other episode", "shared-voice"),
        ]
        for first_duration, next_start, last_start in (
            (9000, 40000, 44000),
            (15000, 40000, 44000),
            (40000, 40000, 44000),
            (40001, 40001, 44000),
            (45000, 45000, 47000),
        ):
            with self.subTest(first_duration=first_duration):
                durations = {2: first_duration, 3: 2000, 4: 500, 5: 60000, 6: 1000}
                placed = {p.row.row_number: p for p in place_dialogue(list(reversed(rows)), durations)}
                self.assertEqual([placed[n].start_ms for n in (2, 3, 4, 5, 6)],
                                 [0, next_start, last_start, 11000, 1000])
                self.assertTrue(all(p.start_ms >= p.row.start_ms for p in placed.values()))
                self.assertEqual(placed[3].end_ms, next_start + 2000)
                self.assertEqual(placed[3].delay_ms, next_start - 40000)

    def test_safe_filenames_remain_unique(self):
        names = track_names({("001", "A/B"), ("001", "A:B"), ("001", "a_b"), ("../../", "...")})
        self.assertEqual(len({name.casefold() for name in names.values()}), 4)
        self.assertTrue(all("/" not in name and "\\" not in name for name in names.values()))
        names = timeline_names({"A/B", "A:B", "a_b", "../../"})
        self.assertEqual(len({name.casefold() for name in names.values()}), 4)
        self.assertTrue(all("/" not in name and "\\" not in name for name in names.values()))

    def test_report_exact_schema_and_frame_rate_calculations(self):
        headers = ["excel_row", "episode", "character", "voice_id", "start_TCR", "end_TCR",
                   "duration", "actual_duration_ms", "overrun_amnt_ms", "status", "track",
                   "adjusted_start_TCR", "adjusted_end_TCR", "fit_speed"]
        expected_ends = {24: "00:00:01:05", 25: "00:00:01:05", 30: "00:00:01:06",
                         50: "00:00:01:10", 60: "00:00:01:11"}
        path = self.workbook([dialogue(start="00:00:00:01", end="00:00:01:01")])
        for fps, end in expected_ends.items():
            with self.subTest(fps=fps):
                output = self.root / f"report-{fps}.zip"
                run_pipeline(path, output, fps=fps, synthesize=lambda *_: wav_bytes(1173), progress=lambda _: None,
                     fit_lines=False)
                with zipfile.ZipFile(output) as archive:
                    records = list(csv.reader(io.StringIO(archive.read("processing_report.csv").decode("utf-8-sig"))))
                    self.assertEqual(records, [headers, [
                        "2", "001", "Alice", "alice-voice", "00:00:00:01", "00:00:01:01",
                        "1000", "1173", "173", "TCR_OUT_OVERRUN", "Episode_001_Alice.wav",
                        "00:00:00:01", end, "1.00",
                    ]])
                    with wave.open(io.BytesIO(archive.read("tracks/Episode_001_Alice.wav"))) as audio:
                        exact_end = timecode_to_ms("00:00:00:01", fps) + 1173
                        self.assertEqual(audio.getnframes(), exact_end * 24)
                        self.assertAlmostEqual(timecode_to_ms(end, fps), exact_end, delta=500 / fps + 0.5)

    def test_tcr_rounding_rollover_and_frame_roundtrip(self):
        for fps in (24, 25, 30, 50, 60):
            with self.subTest(fps=fps):
                for frame in range(fps):
                    original = f"01:02:03:{frame:02d}"
                    self.assertEqual(format_tcr(timecode_to_ms(original, fps), fps), original)
                self.assertEqual(format_tcr(999, fps), "00:00:01:00")
                self.assertEqual(format_tcr(59999, fps), "00:01:00:00")
                self.assertEqual(format_tcr(3599999, fps), "01:00:00:00")
        self.assertEqual(format_tcr(20, 25), "00:00:00:01")
        self.assertEqual(format_tcr(19, 25), "00:00:00:00")

    def test_failed_generation_does_not_publish_partial_archive(self):
        path = self.workbook([dialogue(), dialogue(text="second")])
        output = self.root / "failed.zip"
        synthesize = Mock(side_effect=[wav_bytes(), RuntimeError("Provider failure")])
        with self.assertRaisesRegex(RuntimeError, "Excel row 3: Provider failure"):
            run_pipeline(path, output, synthesize=synthesize, progress=lambda _: None)
        self.assertFalse(output.exists())
        self.assertEqual(list(self.root.glob("ad-v2-*")), [])

    def test_existing_output_is_not_overwritten(self):
        path = self.workbook([dialogue()])
        output = self.root / "existing.zip"
        output.write_bytes(b"existing")
        synthesize = Mock()
        with self.assertRaisesRegex(ValueError, "already exists"):
            run_pipeline(path, output, synthesize=synthesize)
        synthesize.assert_not_called()
        self.assertEqual(output.read_bytes(), b"existing")

    def test_web_upload_status_download_and_template(self):
        app = create_app({"TESTING": True, "RUNTIME_DIR": self.root / "runtime",
                          "SYNTHESIZE": lambda *_: wav_bytes()})
        client = app.test_client()
        self.assertEqual(client.get("/").status_code, 200)
        for asset in ("style.css", "app.js", "zibanka_logo.jpeg"):
            response = client.get(f"/static/{asset}")
            self.assertEqual(response.status_code, 200)
            response.close()
        template = client.get("/template")
        book = load_workbook(io.BytesIO(template.data))
        self.assertEqual(tuple(cell.value for cell in book.active[1]), TEMPLATE_HEADERS)
        book.close()
        path = self.workbook([dialogue()])
        with patch("v2.app.threading.Thread") as thread:
            thread.side_effect = lambda *, target, daemon: Mock(start=target)
            response = client.post("/upload", data={"file": (io.BytesIO(path.read_bytes()), "input.xlsx"),
                                                   "fps": "25", "output_name": "English dialogue.zip"})
        self.assertEqual(response.status_code, 202)
        job_id = response.json["job_id"]
        status = client.get(f"/status/{job_id}").json
        self.assertEqual(status["status"], "done")
        self.assertEqual(status["percent"], 100)
        self.assertEqual(status["download_name"], "English dialogue.zip")
        response = client.get(f"/download/{job_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="English dialogue.zip"', response.headers["Content-Disposition"])
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertIn("tracks/Episode_001_Alice.wav", archive.namelist())
            self.assertIn("lines/001_Alice_2.wav", archive.namelist())
            self.assertIn("Episode_001_timeline.csv", archive.namelist())
            self.assertIn("script.xlsx", archive.namelist())
        response.close()
        # run.json lets the Dubbing Quality Check list and import finished runs.
        record = json.loads((self.root / "runtime" / job_id / "run.json").read_text())
        self.assertEqual((record["status"], record["fps"], record["episodes"]), ("done", 25, ["001"]))
        self.assertEqual(record["download_name"], "English dialogue.zip")
        self.assertEqual(client.get("/download/unknown").status_code, 404)

    def test_web_rejects_invalid_workbook_without_starting_generation(self):
        synthesize = Mock()
        app = create_app({"TESTING": True, "RUNTIME_DIR": self.root / "runtime", "SYNTHESIZE": synthesize})
        client = app.test_client()
        path = self.workbook([dialogue(character="")])
        response = client.post("/upload", data={"file": (io.BytesIO(path.read_bytes()), "input.xlsx")})
        self.assertEqual(response.status_code, 400)
        self.assertIn("CHARACTERS", response.json["error"])
        synthesize.assert_not_called()
        response = client.post("/upload", data={"file": (io.BytesIO(b"broken"), "input.xlsx")})
        self.assertEqual(response.status_code, 400)

    # ---- pass 2: fit each line to its slot with FFmpeg --------------------------------------------
    @staticmethod
    def spoken_clip(speech_ms, lead_ms=150, tail_ms=200, freq=440):
        """Speech surrounded by silence, as ElevenLabs returns it."""
        speech = Sine(freq, sample_rate=24000).to_audio_segment(duration=speech_ms).set_channels(1)
        clip = (AudioSegment.silent(lead_ms, frame_rate=24000) + speech
                + AudioSegment.silent(tail_ms, frame_rate=24000))
        return clip.set_sample_width(2).export(io.BytesIO(), format="wav").getvalue()

    def test_fit_speed_measures_speech_and_holds_limits(self):
        self.assertEqual(speech_span(self.spoken_clip(1000)), (150, 1150, 1350))
        self.assertEqual(fit_speed(self.spoken_clip(2600), 2000), 1.3)              # too long: faster
        self.assertEqual(fit_speed(self.spoken_clip(800), 1000), 0.8)               # too short: slower
        self.assertEqual(fit_speed(self.spoken_clip(1000 + FIT_TOLERANCE_MS), 1000), 1.0)  # within tolerance
        self.assertEqual(fit_speed(self.spoken_clip(4000), 1000), MAX_FIT_SPEED)    # held at 1.5x
        self.assertEqual(fit_speed(self.spoken_clip(200), 1000), MIN_FIT_SPEED)     # held at 0.7x
        self.assertEqual(fit_speed(wav_bytes(500, value=0), 1000), 1.0)             # silence: left alone
        clip = self.spoken_clip(500)
        with patch("v2.audio_speed.shutil.which", return_value=None):
            self.assertIs(change_speed(clip, 1.0), clip)                            # 1.00x needs no FFmpeg
            with self.assertRaisesRegex(RuntimeError, "FFmpeg"):
                change_speed(clip, 1.2)
        for bad in ("", "bad", "nan", 0.69, 1.51):
            with self.assertRaises(ValueError):
                change_speed(clip, bad)

    def test_needed_speed_and_retime_beyond_one_atempo_stage(self):
        # Unlimited speed needed (Quality Check's suggestion), also for a clip already sped up.
        self.assertAlmostEqual(needed_speed(self.spoken_clip(4000), 1000), 4.0, places=2)
        self.assertAlmostEqual(needed_speed(self.spoken_clip(1000), 2000, clip_speed=1.5), 0.75, places=2)
        # Quality Check takes a 1.50x clip to 0.70x: a ratio below FFmpeg's 0.5 per atempo stage.
        clip = self.spoken_clip(1000, lead_ms=0, tail_ms=0)
        for ratio in (0.7 / 1.5, 1.5 / 0.7):
            slowed = AudioSegment.from_wav(io.BytesIO(retime(clip, ratio)))
            self.assertAlmostEqual(len(slowed), 1000 / ratio, delta=40)
        self.assertIs(retime(clip, 1.0), clip)
        with self.assertRaises(ValueError):
            retime(clip, 0)

    def test_pipeline_fits_lines_and_outputs_tracks_and_lines(self):
        path = self.workbook([
            dialogue(start="00:00:00:00", end="00:00:02:00", text="long"),       # 2.6 s speech in 2.0 s
            dialogue(start="00:00:03:00", end="00:00:04:00", text="short"),      # 0.8 s speech in 1.0 s
            dialogue(start="00:00:05:00", end="00:00:06:00", text="fits"),       # 1.05 s: within 0.1 s
            dialogue(character="Bob", voice="b", start="00:00:01:00", end="00:00:02:00", text="huge"),
        ])
        speech = {"long": 2600, "short": 800, "fits": 1050, "huge": 4000}
        output = self.root / "fitted.zip"
        summary = run_pipeline(path, output, synthesize=lambda text, _v: self.spoken_clip(speech[text]),
                               progress=lambda _: None)
        self.assertEqual(summary["fitted_rows"], 3)
        with zipfile.ZipFile(output) as archive:
            names = set(archive.namelist())
            lines = sorted(n for n in names if n.startswith("lines/"))
            self.assertEqual(lines, ["lines/001_Alice_2.wav", "lines/001_Alice_3.wav", "lines/001_Alice_4.wav",
                                     "lines/001_Bob_5.wav"])              # one per dubbed row
            self.assertEqual(sorted(n for n in names if n.startswith("tracks/")),
                             ["tracks/Episode_001_Alice.wav", "tracks/Episode_001_Bob.wav"])
            report = {int(r["excel_row"]): r for r in csv.DictReader(io.StringIO(
                archive.read("processing_report.csv").decode("utf-8-sig")))}
            self.assertEqual({row: r["fit_speed"] for row, r in report.items()},
                             {2: "1.30", 3: "0.80", 4: "1.00", 5: "1.50"})
            # Fitted lines now speak for (about) their slot; the held line is as fast as allowed.
            for row, slot in ((2, 2000), (3, 1000)):
                start, end, _ = speech_span(archive.read(f"lines/001_Alice_{row}.wav"))
                self.assertAlmostEqual(end - start, slot, delta=40)
            start, end, _ = speech_span(archive.read("lines/001_Bob_5.wav"))
            self.assertAlmostEqual(end - start, 4000 / 1.5, delta=40)
            # A line within the tolerance is the ElevenLabs take, untouched.
            self.assertEqual(archive.read("lines/001_Alice_4.wav"), self.spoken_clip(1050))
            # Tracks are built from exactly the clips in lines/, at their placements.
            alice = AudioSegment.from_wav(io.BytesIO(archive.read("tracks/Episode_001_Alice.wav")))
            line = AudioSegment.from_wav(io.BytesIO(archive.read("lines/001_Alice_3.wav")))
            start_ms = timecode_to_ms(report[3]["adjusted_start_TCR"])
            self.assertEqual(alice[start_ms:start_ms + len(line)].raw_data, line.raw_data)
            self.assertEqual(int(report[3]["actual_duration_ms"]), len(line))

    def test_missing_ffmpeg_is_refused_before_any_generation(self):
        synthesize = Mock()
        runtime = self.root / "runtime"
        app = create_app({"TESTING": True, "RUNTIME_DIR": runtime, "SYNTHESIZE": synthesize})
        path = self.workbook([dialogue()])
        with patch("v2.audio_speed.shutil.which", return_value=None), patch("v2.app.threading.Thread") as thread:
            response = app.test_client().post("/upload", data={
                "file": (io.BytesIO(path.read_bytes()), "input.xlsx")})
            self.assertEqual(response.status_code, 503)
            self.assertIn("FFmpeg", response.json["error"])
            with self.assertRaisesRegex(RuntimeError, "FFmpeg"):
                run_pipeline(path, self.root / "missing.zip", synthesize=synthesize)
            thread.assert_not_called()
        synthesize.assert_not_called()
        self.assertFalse(runtime.exists())
        self.assertNotIn('name="speed"', app.test_client().get("/").get_data(as_text=True))

    def test_fit_failure_discards_output(self):
        path = self.workbook([dialogue()])
        output = self.root / "failed-fit.zip"
        with patch("v2.pipeline.fit_to_slot", side_effect=RuntimeError("Audio speed adjustment timed out")):
            with self.assertRaisesRegex(RuntimeError, "Excel row 2: Audio speed adjustment timed out"):
                run_pipeline(path, output, synthesize=lambda *_: wav_bytes(), progress=lambda _: None)
        self.assertFalse(output.exists())
        self.assertEqual(list(self.root.glob("ad-v2-*")), [])

    def test_blank_voice_rows_are_skipped_before_validation_and_generation(self):
        path = self.workbook([
            dialogue(voice=None, start="invalid", text="Skip empty cell"),
            dialogue(voice="", character="", text="Skip empty string"),
            dialogue(voice=" \t ", text="Skip whitespace"),
            dialogue(text="Generate this"),
        ])
        rows = parse_excel(path)
        self.assertEqual([row.row_number for row in rows], [5])
        synthesize = Mock(return_value=wav_bytes())
        summary = run_pipeline(path, self.root / "skipped.zip", synthesize=synthesize,
                               progress=lambda _: None)
        synthesize.assert_called_once_with("Generate this", "alice-voice")
        self.assertEqual(summary["rows"], 1)
        self.assertEqual(summary["tracks"], 1)

    def test_all_blank_voices_produce_no_api_calls(self):
        path = self.workbook([dialogue(voice=None), dialogue(voice="  ")])
        synthesize = Mock()
        with self.assertRaisesRegex(WorkbookValidationError, "No dialogue rows found"):
            run_pipeline(path, self.root / "empty.zip", synthesize=synthesize)
        synthesize.assert_not_called()

    def test_elevenlabs_request_uses_row_voice_and_converts_pcm(self):
        response = Mock()
        response.read.return_value = b"\x00\x00" * 240
        response.__enter__ = Mock(return_value=response)
        response.__exit__ = Mock(return_value=False)
        with patch("v2.tts_client.urlopen", return_value=response) as request, \
                patch.dict("os.environ", {"AD_V2_MODEL_ID": "ignored-legacy-override"}):
            result = ElevenLabsClient(api_key="test-key")("English dialogue", "voice/id")
        sent = request.call_args.args[0]
        self.assertIn("voice%2Fid?output_format=pcm_24000", sent.full_url)
        self.assertEqual(json.loads(sent.data), {
            "text": "English dialogue", "model_id": MODEL_ID,
            "voice_settings": {"stability": 0.5, "similarity_boost": 0.75},
        })
        with wave.open(io.BytesIO(result)) as wav:
            self.assertEqual(wav.getnframes(), 240)
            self.assertEqual(wav.getframerate(), 24000)

    def test_tag_only_rows_get_punctuation_without_added_dialogue(self):
        samples = (
            ("[Gasps]", "[Gasps] ..."),
            ("[LAUGHING]", "[LAUGHING] ..."),
            ("[sighs] [laughs]", "[sighs] [laughs] ..."),
            ("[Gasps] Oh no.", "[Gasps] Oh no."),
        )
        for index, (text, expected) in enumerate(samples):
            with self.subTest(text=text):
                path = self.workbook([dialogue(text=text)])
                response = Mock()
                response.read.return_value = b"\x00\x00" * 240
                response.__enter__ = Mock(return_value=response)
                response.__exit__ = Mock(return_value=False)
                with patch("v2.tts_client.urlopen", return_value=response) as request:
                    run_pipeline(path, self.root / f"sample-{index}.zip",
                                 synthesize=ElevenLabsClient(api_key="test-key"),
                                 progress=lambda _: None)
                self.assertEqual(json.loads(request.call_args.args[0].data)["text"], expected)

    def test_tag_punctuation_is_idempotent_and_preserves_spoken_text(self):
        for text in ("Hello.", "[Gasps] Oh no.", "[laughing] ...", "[Unclosed tag", ""):
            with self.subTest(text=text):
                self.assertEqual(prepare_speech_text(text), text)
        self.assertEqual(prepare_speech_text("  [LAUGHING] \n"), "  [LAUGHING] ...")
        prepared = prepare_speech_text("[Gasps]")
        self.assertEqual(prepare_speech_text(prepared), prepared)

    def test_provider_error_reason_is_reported_with_row_and_credentials_redacted(self):
        path = self.workbook([dialogue(text="[Gasps]")])
        body = json.dumps({"detail": {"status": "example_provider_code",
                                     "message": "Example reason test-key"}}).encode()
        error = HTTPError("https://api.elevenlabs.io", 400, "Bad Request", {}, io.BytesIO(body))
        with patch("v2.tts_client.urlopen", side_effect=error) as request:
            with self.assertRaises(RuntimeError) as caught:
                run_pipeline(path, self.root / "error.zip",
                             synthesize=ElevenLabsClient(api_key="test-key"),
                             progress=lambda _: None)
        message = str(caught.exception)
        self.assertIn("Excel row 2", message)
        self.assertIn("HTTP 400", message)
        self.assertIn("example_provider_code: Example reason [redacted]", message)
        self.assertNotIn("test-key", message)
        self.assertEqual(request.call_count, 1)

    def test_non_json_error_response_is_not_reflected(self):
        error = HTTPError("https://api.elevenlabs.io", 502, "Bad Gateway", {},
                          io.BytesIO(b"<html>Unexpected proxy response</html>"))
        with patch("v2.tts_client.urlopen", side_effect=error):
            with self.assertRaisesRegex(RuntimeError, "HTTP 502.*no structured error details"):
                ElevenLabsClient(api_key="test-key")("[Gasps]", "voice-id")


if __name__ == "__main__":
    unittest.main()
