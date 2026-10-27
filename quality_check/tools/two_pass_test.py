"""Two-pass speed test: how well does speeding a line up (or down) make it fit its TCR slot?

Pass 1 dubs every line exactly as AD version 2 does (eleven_v4, same voice settings, no speed).
We measure where the speech starts and ends inside each clip (ignoring the silence ElevenLabs adds
around it) and compare that spoken length with the slot (TCR OUT - TCR IN). For every line more than
0.10 s off, pass 2 changes the speed of the pass 1 take by spoken / slot:

* FFmpeg (default): the `atempo` filter on the same take, pitch preserved: AD version 2's pass 2
  (``v2/audio_speed.py``), now built into the tools. No new ElevenLabs request. Held between 0.7x and 1.5x.
* ElevenLabs (only with --compare): the line dubbed again with ElevenLabs' `speed` setting, held
  between 0.7x and 1.2x (ElevenLabs' range), so both methods can be heard side by side.

Each line is measured on its own at its TCR IN (AD version 2's "wait for the previous line" rule is
left out on purpose). Results go to an Excel workbook plus a folder of clips.

Run from the repository root (uses the ElevenLabs key from quality_check/.env or v2/.env):

    python -m quality_check.tools.two_pass_test EP01_script.xlsx
    python -m quality_check.tools.two_pass_test EP01_script.xlsx --compare

Results are saved in quality_check/tools_outputs/: two_pass_results_<date>_<time>.xlsx and a matching
_audio folder, so earlier runs are never overwritten. --out NAME.xlsx picks the file name instead.
"""

import argparse
from datetime import datetime
import io
import json
import os
from pathlib import Path
import shutil
import sys
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen
import wave

from dotenv import dotenv_values
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from v2.audio_speed import FIT_TOLERANCE_MS, MAX_FIT_SPEED, MIN_FIT_SPEED, change_speed, speech_span
from v2.excel_parser import SUPPORTED_FPS, parse_excel
from v2.exporter import format_tcr
from v2.tts_client import MODEL_ID, SAMPLE_RATE, VOICE_SETTINGS, prepare_speech_text, provider_error_detail

FFMPEG_MIN, FFMPEG_MAX = MIN_FIT_SPEED, MAX_FIT_SPEED   # AD version 2's pass 2 limits (0.7x, 1.5x)
ELEVENLABS_MIN, ELEVENLABS_MAX = 0.7, 1.2      # ElevenLabs' accepted speed range
TOLERANCE_S = FIT_TOLERANCE_MS / 1000          # a line within this of its slot is left as it is

HERE = Path(__file__).resolve().parents[1]
OUTPUT_DIR = HERE / "tools_outputs"     # every run's workbook and audio folder go here


# ---------------------------------------------------------------------------------- ElevenLabs

class ElevenLabsWithSpeed:
    """AD version 2's request (same endpoint, model, voice settings and PCM output), plus an
    optional `speed` in voice_settings. Pass 1 sends no speed, exactly like the live dub."""

    def __init__(self, api_key):
        self.api_key = api_key

    def __call__(self, text, voice_id, speed=None):
        settings = dict(VOICE_SETTINGS)
        if speed is not None:
            settings["speed"] = speed
        request = Request(
            f"https://api.elevenlabs.io/v1/text-to-speech/{quote(voice_id, safe='')}?output_format=pcm_{SAMPLE_RATE}",
            data=json.dumps({"text": prepare_speech_text(text), "model_id": MODEL_ID,
                             "voice_settings": settings}).encode("utf-8"),
            headers={"Content-Type": "application/json", "Accept": "application/octet-stream",
                     "xi-api-key": self.api_key}, method="POST")
        try:
            with urlopen(request, timeout=120) as response:
                pcm = response.read()
        except HTTPError as exc:
            detail = provider_error_detail(exc, self.api_key)
            raise RuntimeError(f"HTTP {exc.code}{': ' + detail if detail else ''}") from exc
        except (URLError, TimeoutError) as exc:
            raise RuntimeError("connection failed or timed out") from exc
        if not pcm or len(pcm) % 2:
            raise RuntimeError("empty or invalid audio returned")
        out = io.BytesIO()
        with wave.open(out, "wb") as wav:
            wav.setnchannels(1)
            wav.setsampwidth(2)
            wav.setframerate(SAMPLE_RATE)
            wav.writeframes(pcm)
        return out.getvalue()


def load_api_key():
    for env_file in (HERE / ".env", HERE.parent / "v2" / ".env"):
        for key, value in dotenv_values(env_file).items():
            if value and not os.environ.get(key):
                os.environ[key] = value
    key = os.environ.get("ELEVENLABS_API_KEY", "").strip()
    if not key:
        sys.exit("No ElevenLabs key: set ELEVENLABS_API_KEY in quality_check/.env or v2/.env.")
    return key


# ---------------------------------------------------------------------------------- measuring

def clamp(needed, low, high):
    return round(min(high, max(low, needed)), 2)


def limit_note(method, needed, low, high):
    if round(needed, 2) > high:
        return f"{method}: needs {needed:.2f}x, held at the {high:.1f}x limit."
    if round(needed, 2) < low:
        return f"{method}: needs {needed:.2f}x, held at the {low:.1f}x limit."
    return ""


# ---------------------------------------------------------------------------------- the test

def run(script, out, *, synthesize, compare=False, fps=25, sheet=None, rows=None, tolerance=TOLERANCE_S,
        log=print):
    script, out = Path(script), Path(out)
    if shutil.which("ffmpeg") is None:
        sys.exit("FFmpeg is not installed or not on PATH; it is needed for pass 2.")
    lines = parse_excel(script, fps=fps, sheet_name=sheet)
    if rows:
        lines = [line for line in lines if line.row_number in rows]
    if not lines:
        sys.exit("No dialogue rows to test.")
    audio_dir = out.with_name(out.stem + "_audio")
    audio_dir.mkdir(parents=True, exist_ok=True)
    chars = {"pass1": 0, "elevenlabs": 0}
    results = []

    for n, row in enumerate(lines, start=1):
        result = {"row": row, "slot_s": (row.end_ms - row.start_ms) / 1000, "pass1": None,
                  "ff_speed": None, "ffmpeg": None, "el_speed": None, "elevenlabs": None, "notes": []}
        results.append(result)
        stem = f"row{row.row_number}_{row.character}"
        log(f"[{n}/{len(lines)}] row {row.row_number} {row.character}: pass 1 (ElevenLabs)")
        try:
            take = synthesize(row.text, row.voice_id, None)
            chars["pass1"] += len(prepare_speech_text(row.text))
        except Exception as exc:
            result["notes"].append(f"Pass 1 failed: {exc}")
            log(f"    {result['notes'][-1]}")
            continue
        (audio_dir / f"{stem}_pass1.wav").write_bytes(take)
        result["pass1"] = speech_span(take)
        start, end, _clip = result["pass1"]
        spoken = (end - start) / 1000
        needed = spoken / result["slot_s"]
        if abs(result["slot_s"] - spoken) <= tolerance:
            result["notes"].append(f"Within {tolerance:.2f} s after pass 1: left as it is.")
            continue

        # Pass 2 with FFmpeg: the same take, time-stretched.
        speed = clamp(needed, FFMPEG_MIN, FFMPEG_MAX)
        result["ff_speed"] = speed
        note = limit_note("FFmpeg", needed, FFMPEG_MIN, FFMPEG_MAX)
        if note:
            result["notes"].append(note)
        log(f"    pass 2 (FFmpeg) at {speed:.2f}x")
        try:
            wav = change_speed(take, speed)  # AD version 2's pass 2: FFmpeg atempo
            (audio_dir / f"{stem}_pass2_ffmpeg_{speed:.2f}x.wav").write_bytes(wav)
            result["ffmpeg"] = speech_span(wav)
        except Exception as exc:
            result["notes"].append(f"FFmpeg failed at {speed:.2f}x: {exc}")

        # Optional comparison: the line dubbed again with ElevenLabs' speed setting.
        if compare:
            speed = clamp(needed, ELEVENLABS_MIN, ELEVENLABS_MAX)
            result["el_speed"] = speed
            note = limit_note("ElevenLabs", needed, ELEVENLABS_MIN, ELEVENLABS_MAX)
            if note:
                result["notes"].append(note)
            log(f"    pass 2 (ElevenLabs) at {speed:.2f}x")
            try:
                wav = synthesize(row.text, row.voice_id, speed)
                chars["elevenlabs"] += len(prepare_speech_text(row.text))
                (audio_dir / f"{stem}_pass2_elevenlabs_{speed:.2f}x.wav").write_bytes(wav)
                result["elevenlabs"] = speech_span(wav)
            except Exception as exc:
                result["notes"].append(f"ElevenLabs refused {speed:.2f}x: {exc}")
                log(f"    {result['notes'][-1]}")

    write_workbook(out, results, compare=compare, fps=fps, tolerance=tolerance, chars=chars,
                   script=script, audio_dir=audio_dir)
    log(f"Done: {out} and {audio_dir}/")
    return results


# ---------------------------------------------------------------------------------- workbook

FONT = "Arial"
SECONDS = '0.00" s"'
SIGNED_SECONDS = '+0.00" s";-0.00" s";0.00" s"'
SPEED = '0.00"x"'
BASE_HEADERS = [
    ("Excel row", 9), ("Character", 16), ("Dialogue", 48), ("TCR IN", 13), ("TCR OUT", 13),
    ("Slot length (s)", 11),
    ("Pass 1 speech start", 13), ("Pass 1 speech end", 13), ("Pass 1 spoken length (s)", 12),
    ("Duration difference (s)", 12), ("Speed needed", 10),
]
PASS2_HEADERS = [("Speed applied", 10), ("speech start", 13), ("speech end", 13), ("spoken length (s)", 12),
                 ("Duration difference after pass 2 (s)", 14), ("How much faster it actually got", 13)]


def write_workbook(out, results, *, compare, fps, tolerance, chars, script, audio_dir):
    book = Workbook()
    ws = book.active
    ws.title = "Lines"
    thin = Side(style="thin", color="D0D5DA")
    head_fill = PatternFill("solid", fgColor="E8EEF4")

    # Column layout: base columns, the FFmpeg pass 2 block, the optional ElevenLabs block, notes.
    blocks = [("FFmpeg", "ffmpeg", "ff_speed", PatternFill("solid", fgColor="F3F8F3"))]
    if compare:
        blocks.append(("ElevenLabs", "elevenlabs", "el_speed", PatternFill("solid", fgColor="F6F2FB")))
    headers = list(BASE_HEADERS)
    block_cols = []
    for label, *_ in blocks:
        block_cols.append(len(headers) + 1)
        names = [f"Speed applied ({label})", f"Pass 2 ({label}) speech start", f"Pass 2 ({label}) speech end",
                 f"Pass 2 ({label}) spoken length (s)", f"Duration difference after pass 2 ({label}) (s)",
                 f"How much faster it actually got ({label})"]
        headers += [(name, width) for name, (_title, width) in zip(names, PASS2_HEADERS)]
    notes_col = len(headers) + 1
    headers.append(("Notes", 46))

    for col, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 54
    ws.freeze_panes = "D2"

    for r, result in enumerate(results, start=2):
        row = result["row"]
        values = {1: row.row_number, 2: row.character, 3: row.text,
                  4: format_tcr(row.start_ms, fps), 5: format_tcr(row.end_ms, fps),
                  6: round(result["slot_s"], 3), notes_col: " ".join(result["notes"])}
        if result["pass1"]:
            start, end, _ = result["pass1"]
            values.update({7: format_tcr(row.start_ms + start, fps), 8: format_tcr(row.start_ms + end, fps),
                           9: round((end - start) / 1000, 3)})
        # Calculated columns are formulas, so the sheet stays right if a number is corrected by hand.
        values[10] = f'=IF(I{r}="","",ROUND(F{r}-I{r},3))'
        values[11] = f'=IF(I{r}="","",ROUND(I{r}/F{r},2))'
        formats = {6: SECONDS, 9: SECONDS, 10: SIGNED_SECONDS, 11: SPEED}
        fills = {}
        for (label, key, speed_key, fill), first in zip(blocks, block_cols):
            sp, st, en, ln, df, ach = range(first, first + 6)
            L = get_column_letter
            values[sp] = result[speed_key]
            if result[key]:
                start, end, _ = result[key]
                values.update({st: format_tcr(row.start_ms + start, fps), en: format_tcr(row.start_ms + end, fps),
                               ln: round((end - start) / 1000, 3)})
            values[df] = f'=IF({L(ln)}{r}="","",ROUND(F{r}-{L(ln)}{r},3))'
            values[ach] = f'=IF({L(ln)}{r}="","",ROUND(I{r}/{L(ln)}{r},2))'
            formats.update({sp: SPEED, ln: SECONDS, df: SIGNED_SECONDS, ach: SPEED})
            fills.update({c: fill for c in range(first, first + 6)})
        for col, value in values.items():
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (3, notes_col))
            if col in formats:
                cell.number_format = formats[col]
            if col in fills:
                cell.fill = fills[col]
    last = len(results) + 1

    # Negative duration difference = speech longer than the slot (red); positive = shorter (green).
    diff_cols = ["J"] + [get_column_letter(first + 4) for first in block_cols]
    for col in diff_cols:
        span = f"{col}2:{col}{last}"
        ws.conditional_formatting.add(span, CellIsRule(operator="lessThan", formula=[f"-{tolerance}"],
                                                       font=Font(name=FONT, size=10, color="B23B33")))
        ws.conditional_formatting.add(span, CellIsRule(operator="greaterThan", formula=[f"{tolerance}"],
                                                       font=Font(name=FONT, size=10, color="1F7A4D")))
    ws.auto_filter.ref = f"A1:{get_column_letter(notes_col)}{last}"

    write_summary(book, results, blocks, block_cols, last, compare=compare, fps=fps, tolerance=tolerance,
                  chars=chars, script=script, audio_dir=audio_dir)
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)


def write_summary(book, results, blocks, block_cols, last, *, compare, fps, tolerance, chars, script, audio_dir):
    sm = book.create_sheet("Summary")
    sm.column_dimensions["A"].width = 56
    sm.column_dimensions["B"].width = 22
    sm.column_dimensions["C"].width = 72
    L = lambda col: f"Lines!{col}2:{col}{last}"  # noqa: E731
    entries = []      # (label, value, note, number format, is input)
    where = {}        # label key -> row number, for formulas that refer to settings

    def add(key, label, value=None, note="", fmt=None, given=False, head=False):
        entries.append((key, label, value, note, fmt, given, head))

    add("h1", "Settings", head=True)
    add("script", "Script", script.name, given=True)
    add("date", "Run on", datetime.now().strftime("%Y-%m-%d %H:%M"), given=True)
    add("model", "ElevenLabs model (pass 1)", MODEL_ID, "Same model and voice settings as AD version 2", given=True)
    add("fps", "Frame rate (fps)", fps, given=True)
    add("tol", "Tolerance (s)", tolerance, "Lines within this of their slot after pass 1 are left as they are",
        SECONDS, given=True)
    add("ffmin", "FFmpeg slowest speed", FFMPEG_MIN, "atempo, pitch preserved, applied to the pass 1 take", SPEED, True)
    add("ffmax", "FFmpeg fastest speed", FFMPEG_MAX, "", SPEED, True)
    if compare:
        add("elmin", "ElevenLabs slowest speed", ELEVENLABS_MIN, "ElevenLabs' accepted range", SPEED, True)
        add("elmax", "ElevenLabs fastest speed", ELEVENLABS_MAX, "", SPEED, True)
    add("detect", "Speech detection", f"{SILENCE_DB} dB below peak",
        "Silence before and after the voice is not counted as speech", given=True)
    add("audio", "Audio clips", audio_dir.name + "/", "Pass 1 and pass 2 WAVs for listening", given=True)
    add("gap", None)
    add("h2", "Results", head=True)
    add("tested", "Lines tested", f"=COUNT({L('A')})")
    add("fit1", "Lines within tolerance after pass 1", f'=COUNTIFS({L("J")},">="&-{{tol}},{L("J")},"<="&{{tol}})')
    add("over", "Lines needing a speed change", f'=COUNT({L(get_column_letter(block_cols[0]))})')
    add("avg1", "Average size of duration difference, pass 1 (s)",
        f'=IFERROR((SUMIF({L("J")},">0")-SUMIF({L("J")},"<0"))/COUNT({L("J")}),"")', "Ignoring sign; all lines", SECONDS)
    for (label, key, _speed_key, _fill), first in zip(blocks, block_cols):
        sp, ln, df, ach = (get_column_letter(first + i) for i in (0, 3, 4, 5))
        low, high = ("ffmin", "ffmax") if key == "ffmpeg" else ("elmin", "elmax")
        add(f"h_{key}", f"Pass 2 with {label}", head=True)
        add(f"fit_{key}", "Lines within tolerance after pass 2",
            f'=COUNTIFS({L("J")},">="&-{{tol}},{L("J")},"<="&{{tol}},{L(sp)},"=")'
            f'+COUNTIFS({L(df)},">="&-{{tol}},{L(df)},"<="&{{tol}})',
            "Lines left as they are count with their pass 1 result")
        add(f"held_{key}", "Lines needing more than the fastest speed", f'=COUNTIF({L("K")},">"&{{{high}}})',
            "Held at the limit; these need shorter text")
        add(f"low_{key}", "Lines needing less than the slowest speed", f'=COUNTIF({L("K")},"<"&{{{low}}})')
        add(f"avg_{key}", "Average size of duration difference after pass 2 (s)",
            f'=IFERROR((SUMIF({L(df)},">0")-SUMIF({L(df)},"<0"))/COUNT({L(df)}),"")', "Ignoring sign; changed lines only",
            SECONDS)
        add(f"sp_{key}", "Average speed applied", f'=IFERROR(AVERAGE({L(sp)}),"")', "", SPEED)
        add(f"ach_{key}", "Average speed actually achieved", f'=IFERROR(AVERAGE({L(ach)}),"")',
            "Pass 1 length / pass 2 length; compare with the speed applied", SPEED)
    add("h3", "ElevenLabs usage", head=True)
    add("c1", "Characters sent, pass 1", chars["pass1"], "Counted by this script; ElevenLabs bills by characters",
        given=True)
    if compare:
        add("c2", "Characters sent, ElevenLabs pass 2", chars["elevenlabs"], given=True)
    add("ct", "Characters sent, total", "=" + "+".join("{%s}" % k for k in (["c1", "c2"] if compare else ["c1"])))
    add("ffnote", "FFmpeg pass 2", "No ElevenLabs request", "Free: it changes the speed of the pass 1 take")

    for r, (key, *_rest) in enumerate(entries, start=1):
        where[key] = r
    for r, (key, label, value, note, fmt, given, head) in enumerate(entries, start=1):
        if label is None:
            continue
        a = sm.cell(row=r, column=1, value=label)
        a.font = Font(name=FONT, size=11 if head else 10, bold=head)
        if head:
            continue
        if isinstance(value, str) and value.startswith("="):
            value = value.format(**{k: f"$B${v}" for k, v in where.items()})
        b = sm.cell(row=r, column=2, value=value)
        # Blue = values written by the script (settings, counts); black = formulas over the Lines sheet.
        b.font = Font(name=FONT, size=10, color="0000FF" if given else "000000")
        b.alignment = Alignment(horizontal="left")
        if fmt:
            b.number_format = fmt
        c = sm.cell(row=r, column=3, value=note)
        c.font = Font(name=FONT, size=9, color="66717D")
    book.move_sheet("Summary", offset=-1)
    book.active = 0


# ---------------------------------------------------------------------------------- command line

def parse_rows(text):
    """'2-10,14' -> {2,...,10,14}"""
    rows = set()
    for part in filter(None, (p.strip() for p in text.split(","))):
        a, _, b = part.partition("-")
        rows.update(range(int(a), int(b or a) + 1))
    return rows


def main():
    parser = argparse.ArgumentParser(description="Two-pass speed test for one dubbing script.")
    parser.add_argument("script", type=Path, help="Dubbing script (.xlsx), same format as AD version 2")
    parser.add_argument("--out", help="Results file name, saved in quality_check/tools_outputs/ "
                                       "(default: two_pass_results_<date>_<time>.xlsx)")
    parser.add_argument("--compare", action="store_true",
                        help="Also dub pass 2 again with ElevenLabs' speed setting, to compare with FFmpeg (uses credits)")
    parser.add_argument("--fps", type=int, choices=SUPPORTED_FPS, default=25)
    parser.add_argument("--sheet", help="Worksheet name (default: first)")
    parser.add_argument("--rows", type=parse_rows, help="Only these Excel rows, e.g. 2-10,14")
    parser.add_argument("--tolerance", type=float, default=TOLERANCE_S, help="Seconds (default 0.10)")
    args = parser.parse_args()
    name = Path(args.out).name if args.out else f"two_pass_results_{datetime.now():%Y%m%d_%H%M}.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    out = OUTPUT_DIR / name
    print(f"Results will be saved in {out.parent}/")
    run(args.script, out, synthesize=ElevenLabsWithSpeed(load_api_key()), compare=args.compare,
        fps=args.fps, sheet=args.sheet, rows=args.rows, tolerance=args.tolerance)


if __name__ == "__main__":
    main()
