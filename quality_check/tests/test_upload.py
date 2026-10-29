"""Run from the repository root: python -m unittest discover -s quality_check/tests -t . -v"""

import io
import json
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import wave
import zipfile

from quality_check.app import create_app
from quality_check.tests.helpers import (ONE_EPISODE, SECOND_EPISODE, dub_zip, fake_voice, have_pydub,
                                         workbook_bytes)
from v2.audio_speed import fit_speed


def make_client(runtime, synthesize=fake_voice, v2_runtime=None):
    return create_app({"RUNTIME_DIR": runtime, "TESTING": True, "SYNTHESIZE": synthesize,
                       "V2_RUNTIME_DIR": v2_runtime or Path(runtime).parent / "v2_runtime"}).test_client()


def natural_ms(text):
    """Length of the fake voice's take for ``text`` (its whole clip is speech)."""
    return max(300, 45 * len(text))


class ScriptCheckTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.runtime = Path(self.tmp.name) / "runtime"
        self.client = make_client(self.runtime)

    def check(self, rows, **form):
        data = {"script": (io.BytesIO(workbook_bytes(rows)), "script.xlsx"), "fps": "25", **form}
        return self.client.post("/check-script", data=data, content_type="multipart/form-data")

    def test_summary_without_saving(self):
        response = self.check(ONE_EPISODE)
        self.assertEqual(response.status_code, 200, response.get_json())
        summary = response.get_json()["summary"]
        self.assertEqual(summary["lines"], 5)
        self.assertEqual([c["name"] for c in summary["characters"]], ["KUHU", "AMRITA"])
        self.assertEqual(summary["cues"][0], [12880, 14360, 0])  # time order, character index
        self.assertFalse(self.runtime.exists())

    def test_multiple_episodes_rejected(self):
        rows = ONE_EPISODE + [["2", 9, "00:00:01:00", "00:00:02:00", "KUHU", "voiceA", "Hi"]]
        response = self.check(rows)
        self.assertEqual(response.status_code, 400)
        self.assertIn("2 episodes", response.get_json()["error"])

    def test_parser_errors_reported(self):
        rows = [["1", 1, "00:00:14:09", "00:00:12:22", "KUHU", "voiceA", "Backwards"]]
        self.assertIn("TCR OUT must be later", self.check(rows).get_json()["error"])

    def test_bad_frame_rate_and_missing_sheet(self):
        self.assertEqual(self.check(ONE_EPISODE, fps="23").status_code, 400)
        self.assertIn("'Nope' was not found", self.check(ONE_EPISODE, sheet="Nope").get_json()["error"])

    def test_corrupt_workbook(self):
        data = {"script": (io.BytesIO(b"not a workbook"), "script.xlsx"), "fps": "25"}
        response = self.client.post("/check-script", data=data, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 400)
        self.assertIn("couldn't be read", response.get_json()["error"])

    def test_missing_api_key_refused_before_anything_is_stored(self):
        with patch.dict(os.environ, {"ELEVENLABS_API_KEY": ""}), \
                patch("quality_check.app.dotenv_values", return_value={}):
            client = make_client(self.runtime, synthesize=None)
            data = {"script": (io.BytesIO(workbook_bytes(ONE_EPISODE)), "script.xlsx"), "fps": "25"}
            response = client.post("/sessions", data=data, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 503)
        self.assertIn("ELEVENLABS_API_KEY", response.get_json()["error"])
        self.assertFalse(self.runtime.exists() and any(self.runtime.iterdir()))


@unittest.skipUnless(have_pydub(), "dubbing uses V2's pydub track builder (Python 3.11/3.12)")
class DubbingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.runtime = Path(self.tmp.name) / "runtime"
        self.calls = []

    def voice(self, text, voice_id):
        self.calls.append(text)
        return fake_voice(text, voice_id)

    def start(self, client, rows=ONE_EPISODE, **form):
        data = {"script": (io.BytesIO(workbook_bytes(rows)), "script.xlsx"), "fps": "25", **form}
        return client.post("/sessions", data=data, content_type="multipart/form-data")

    def wait(self, client, sid, timeout=30):
        deadline = time.time() + timeout
        while time.time() < deadline:
            record = client.get(f"/session/{sid}").get_json()
            if record["status"] != "dubbing":
                return record
            time.sleep(0.05)
        self.fail("dubbing did not finish")

    def attach_video(self, client, sid):
        data = {"video": (io.BytesIO(b"fake-video-bytes"), "episode.mp4")}
        return client.post(f"/session/{sid}/video", data=data, content_type="multipart/form-data")

    def test_tracks_are_identical_to_ad_version_2(self):
        client = make_client(self.runtime, self.voice)
        sid = self.start(client).get_json()["session_id"]
        record = self.wait(client, sid)
        self.assertEqual(record["status"], "dubbed", record.get("error"))
        self.assertEqual(len(self.calls), 5)

        # The same script through V2's own pipeline must give byte-identical audio per character.
        expected = {}
        with zipfile.ZipFile(io.BytesIO(dub_zip(self.tmp.name, ONE_EPISODE))) as archive:
            for name in ("KUHU", "AMRITA"):
                with wave.open(io.BytesIO(archive.read(f"pass2/tracks/ep1_{name.lower()}.wav"))) as wav:
                    expected[name] = wav.readframes(wav.getnframes())
        folder = self.runtime / sid
        for track in record["audio"]["tracks"]:
            with wave.open(str(folder / "tracks" / f"{track['index']}.wav")) as wav:
                self.assertEqual(wav.readframes(wav.getnframes()), expected[track["character"]], track["character"])

        lines = {line["row"]: line for line in record["audio"]["lines"]}
        self.assertGreater(lines[4]["overrun_ms"], 0)                         # KUHU's long line...
        self.assertEqual(lines[5]["audio_start_ms"], lines[4]["audio_end_ms"])  # ...pushes the next one
        self.assertEqual(len(list((folder / "clips").glob("*.take.wav"))), 5)   # the takes as dubbed
        self.assertEqual(len(list((folder / "clips").glob("*.wav"))), 10)       # plus what plays
        self.assertTrue(any("late" in w for w in record["audio"]["warnings"]))

    def test_review_is_ready_only_with_both_dub_and_video(self):
        gate = threading.Event()

        def slow_voice(text, voice_id):
            gate.wait(5)
            return fake_voice(text, voice_id)

        client = make_client(self.runtime, slow_voice)
        sid = self.start(client).get_json()["session_id"]
        record = self.attach_video(client, sid).get_json()      # upload while dubbing
        self.assertEqual((record["status"], record["ready"]), ("dubbing", False))
        self.assertEqual(record["progress"]["total"], 5)
        gate.set()
        record = self.wait(client, sid)
        self.assertTrue(record["ready"])
        response = client.get(record["video_url"], headers={"Range": "bytes=0-3"})
        self.assertEqual((response.status_code, response.data), (206, b"fake"))
        response.close()
        track = client.get(record["audio"]["tracks"][0]["url"], headers={"Range": "bytes=0-3"})
        self.assertEqual((track.status_code, track.data), (206, b"RIFF"))
        track.close()
        self.assertEqual(set(client.get(record["peaks_url"]).get_json()["tracks"]), {"0", "1"})

    def test_failed_line_reported_and_retry_resumes_without_regenerating(self):
        fail = {"on": True}

        def flaky(text, voice_id):
            if fail["on"] and text.startswith("[Gasps]"):
                raise RuntimeError("ElevenLabs returned HTTP 429 (model eleven_v4): too_many_requests")
            return self.voice(text, voice_id)

        client = make_client(self.runtime, flaky)
        sid = self.start(client).get_json()["session_id"]
        record = self.wait(client, sid)
        self.assertEqual(record["status"], "error")
        self.assertEqual(record["failed_row"], 4)
        self.assertIn("Row 4 (KUHU at 00:00:30:00)", record["error"])
        self.assertIn("429", record["error"])
        before = len(self.calls)                     # rows 2 and 3 were already generated

        fail["on"] = False
        self.assertEqual(client.post(f"/session/{sid}/retry").status_code, 200)
        record = self.wait(client, sid)
        self.assertEqual(record["status"], "dubbed")
        self.assertEqual(len(self.calls) - before, 3)  # only rows 4, 5 and 6: nothing paid twice

    def test_delete_stops_dubbing(self):
        started = threading.Event()

        def slow_voice(text, voice_id):
            started.set()
            time.sleep(0.2)
            return self.voice(text, voice_id)

        client = make_client(self.runtime, slow_voice)
        sid = self.start(client).get_json()["session_id"]
        started.wait(5)
        self.assertEqual(client.delete(f"/session/{sid}").status_code, 204)
        time.sleep(0.5)
        self.assertLessEqual(len(self.calls), 2)     # stopped before the next request
        self.assertEqual(client.get(f"/session/{sid}").status_code, 404)

    def test_every_line_is_fitted_to_its_slot(self):
        # Pass 2 as in AD version 2: the take is ElevenLabs as delivered, what plays is fitted.
        client = make_client(self.runtime, self.voice)
        sid = self.start(client).get_json()["session_id"]
        record = self.wait(client, sid)
        self.assertNotIn("speed", record)                                   # no speed dial any more
        folder = self.runtime / sid
        lines = {line["row"]: line for line in record["audio"]["lines"]}
        self.assertEqual({row: line["fit_speed"] for row, line in lines.items()},
                         {2: 0.7, 3: 0.7, 4: 1.5, 5: 0.7, 6: 0.81})
        self.assertEqual(lines[4]["needed_speed"], 2.16)                    # held at the 1.5x limit
        for row, line in lines.items():
            take = (folder / "clips" / f"{row}.take.wav").read_bytes()
            self.assertEqual(fit_speed(take, line["end_ms"] - line["start_ms"]), line["fit_speed"])
            self.assertEqual(line["line_speed"], line["fit_speed"])
            self.assertFalse(line["speed_set"])
            self.assertAlmostEqual(line["audio_end_ms"] - line["audio_start_ms"],
                                   natural_ms(line["text"]) / line["fit_speed"], delta=30)

    def test_generating_needs_ffmpeg(self):
        client = make_client(self.runtime, self.voice)
        with patch("v2.audio_speed.shutil.which", return_value=None):
            response = self.start(client)
        self.assertEqual(response.status_code, 503)
        self.assertIn("FFmpeg", response.get_json()["error"])
        self.assertEqual(self.calls, [])

    # -- Block 5: regenerate one line ---------------------------------------------------------
    def dubbed(self, voice=None):
        client = make_client(self.runtime, voice or self.voice)
        sid = self.start(client).get_json()["session_id"]
        record = self.wait(client, sid)
        self.assertEqual(record["status"], "dubbed")
        return client, sid, record

    def regenerate(self, client, sid, row, text):
        return client.post(f"/session/{sid}/lines/{row}", json={"text": text})

    def track_bytes(self, sid, index):
        return (self.runtime / sid / "tracks" / f"{index}.wav").read_bytes()

    def test_shorter_line_rebuilds_only_that_character(self):
        # AMRITA's row 3 overruns into her row 6 (23:00), which the dub therefore delays.
        client, sid, record = self.dubbed()
        kuhu_before = self.track_bytes(sid, 0)
        lines = {l["row"]: l for l in record["audio"]["lines"]}
        self.assertGreater(lines[6]["delay_ms"], 0)
        self.calls.clear()
        response = self.regenerate(client, sid, 3, "Out!")
        self.assertEqual(response.status_code, 200, response.get_json())
        body = response.get_json()
        self.assertEqual(self.calls, ["Out!"])                        # exactly one ElevenLabs request
        self.assertEqual(body["changed_tracks"], [1])                 # AMRITA only
        self.assertEqual(self.track_bytes(sid, 0), kuhu_before)       # KUHU's file untouched
        new = {l["row"]: l for l in body["audio"]["lines"]}
        self.assertEqual(new[3]["fit_speed"], 0.7)                    # the new take is fitted too
        self.assertAlmostEqual(new[3]["audio_end_ms"] - new[3]["audio_start_ms"], natural_ms("Out!") / 0.7, delta=30)
        self.assertTrue(new[3]["edited"])
        self.assertEqual(new[3]["original"], lines[3]["text"])
        # Row 6 no longer waits for row 3's overrun: it moves back to its own TCR IN.
        self.assertEqual(new[6]["audio_start_ms"], new[6]["start_ms"])
        self.assertEqual(body["moved"], [{"row": 6, "character": "AMRITA",
                                          "shift_ms": new[6]["start_ms"] - lines[6]["audio_start_ms"]}])
        # The rebuilt track holds the new clip at the right place.
        with wave.open(str(self.runtime / sid / "tracks" / "1.wav")) as track, \
                wave.open(str(self.runtime / sid / "clips" / "3.wav")) as clip:
            track.setpos(new[3]["audio_start_ms"] * 24)
            self.assertEqual(track.readframes(500), clip.readframes(500))

    def test_change_to_episode_end_rebuilds_every_track(self):
        # A long take for KUHU's row 5 makes it the last audio in the episode (past 32 s); shortening
        # row 4 then pulls it earlier, so the episode gets shorter and (V2 rule: one length for all
        # tracks) every track is rebuilt.
        client, sid, _record = self.dubbed()
        record = self.regenerate(client, sid, 5, "Stop! " * 10).get_json()
        self.assertGreater(record["audio"]["duration_ms"], 32000)
        body = self.regenerate(client, sid, 4, "[Gasps] Too long.").get_json()
        self.assertEqual(sorted(body["changed_tracks"]), [0, 1])
        self.assertLess(body["audio"]["duration_ms"], record["audio"]["duration_ms"])
        self.assertEqual(len({len(self.track_bytes(sid, i)) for i in (0, 1)}), 1)

    def test_back_to_original_text_clears_the_edit(self):
        client, sid, record = self.dubbed()
        original = next(l["text"] for l in record["audio"]["lines"] if l["row"] == 2)
        self.regenerate(client, sid, 2, "Edited words")
        body = self.regenerate(client, sid, 2, original).get_json()
        self.assertFalse(next(l for l in body["audio"]["lines"] if l["row"] == 2)["edited"])
        self.assertEqual(client.get(f"/session/{sid}").get_json()["edits"], {})

    def test_line_past_the_episode_end_rebuilds_all_tracks_to_one_length(self):
        client, sid, record = self.dubbed()
        body = self.regenerate(client, sid, 6, "Stop! " * 60).get_json()  # past the last TCR OUT even at 1.5x
        self.assertEqual(sorted(body["changed_tracks"]), [0, 1])
        self.assertGreater(body["audio"]["duration_ms"], record["audio"]["duration_ms"])
        lengths = {len(self.track_bytes(sid, i)) for i in (0, 1)}
        self.assertEqual(len(lengths), 1)

    def test_failed_regeneration_changes_nothing(self):
        client, sid, record = self.dubbed()
        before = (self.track_bytes(sid, 0), (self.runtime / sid / "clips" / "4.wav").read_bytes())
        app = client.application
        app.config["SYNTHESIZE"] = lambda text, voice: (_ for _ in ()).throw(RuntimeError("HTTP 401: invalid_api_key"))
        response = self.regenerate(client, sid, 4, "New words")
        self.assertEqual(response.status_code, 502)
        self.assertIn("Row 4 (KUHU at 00:00:30:00)", response.get_json()["error"])
        self.assertEqual(before, (self.track_bytes(sid, 0), (self.runtime / sid / "clips" / "4.wav").read_bytes()))
        self.assertEqual(client.get(f"/session/{sid}").get_json()["edits"], {})

    def test_regenerate_validation(self):
        gate = threading.Event()
        client = make_client(self.runtime, lambda t, v: (gate.wait(5), fake_voice(t, v))[1])
        sid = self.start(client).get_json()["session_id"]
        self.assertEqual(self.regenerate(client, sid, 2, "Hi").status_code, 409)   # still dubbing
        gate.set()
        self.wait(client, sid)
        self.assertEqual(self.regenerate(client, sid, 2, "   ").status_code, 400)
        self.assertEqual(self.regenerate(client, sid, 2, "x" * 1501).status_code, 400)
        self.assertEqual(self.regenerate(client, sid, 99, "Hi").status_code, 404)

    # -- line speed ----------------------------------------------------------------------------
    def speed(self, client, sid, row, value):
        return client.post(f"/session/{sid}/lines/{row}/speed", json={"speed": value})

    def test_line_speed_from_the_take_keeps_the_start_and_restores_exactly(self):
        client, sid, record = self.dubbed()
        folder = self.runtime / sid
        take = (folder / "clips" / "3.take.wav").read_bytes()
        fitted = (folder / "clips" / "3.wav").read_bytes()
        kuhu_before = self.track_bytes(sid, 0)
        before = {l["row"]: l for l in record["audio"]["lines"]}
        self.assertEqual(before[3]["fit_speed"], 0.7)
        natural = natural_ms(before[3]["text"])
        self.calls.clear()

        body = self.speed(client, sid, 3, 1.2).get_json()
        self.assertEqual(self.calls, [])                              # no ElevenLabs request
        self.assertEqual(body["changed_tracks"], [1])                 # AMRITA only
        self.assertEqual(self.track_bytes(sid, 0), kuhu_before)
        lines = {l["row"]: l for l in body["audio"]["lines"]}
        self.assertEqual((lines[3]["line_speed"], lines[3]["speed_set"]), (1.2, True))
        self.assertEqual(lines[3]["fit_speed"], 0.7)                  # the suggestion stays pass 2's
        self.assertEqual(lines[3]["audio_start_ms"], before[3]["audio_start_ms"])      # same start
        self.assertAlmostEqual(lines[3]["audio_end_ms"] - lines[3]["audio_start_ms"], natural / 1.2, delta=30)
        self.assertLess(lines[6]["audio_start_ms"], before[6]["audio_start_ms"])       # row 6 waits less
        self.assertEqual([m["row"] for m in body["moved"]], [6])

        # Changes never stack: speeds are measured from the take; 1.00x is the take byte for byte
        # and the fit speed gives the fitted clip back (no longer a reviewer change).
        body = self.speed(client, sid, 3, 1.1).get_json()
        line = next(l for l in body["audio"]["lines"] if l["row"] == 3)
        self.assertAlmostEqual(line["audio_end_ms"] - line["audio_start_ms"], natural / 1.1, delta=30)
        self.speed(client, sid, 3, 1.0)
        self.assertEqual((folder / "clips" / "3.wav").read_bytes(), take)
        body = self.speed(client, sid, 3, 0.7).get_json()
        self.assertEqual((folder / "clips" / "3.wav").read_bytes(), fitted)
        self.assertEqual(client.get(f"/session/{sid}").get_json()["line_speeds"], {})
        self.assertEqual({l["row"]: l["audio_start_ms"] for l in body["audio"]["lines"]},
                         {r: l["audio_start_ms"] for r, l in before.items()})

    def test_regenerating_resets_line_speed(self):
        client, sid, _record = self.dubbed()
        self.speed(client, sid, 3, 1.15)
        body = self.regenerate(client, sid, 3, "[Angry] Get out!").get_json()
        line = next(l for l in body["audio"]["lines"] if l["row"] == 3)
        self.assertEqual((line["line_speed"], line["speed_set"]), (line["fit_speed"], False))
        self.assertEqual(client.get(f"/session/{sid}").get_json()["line_speeds"], {})

    def test_line_speed_validation(self):
        client, sid, _record = self.dubbed()
        self.assertEqual(self.speed(client, sid, 3, 1.5).status_code, 200)    # the limits are allowed
        self.assertEqual(self.speed(client, sid, 3, 0.7).status_code, 200)
        self.assertEqual(self.speed(client, sid, 3, 1.51).status_code, 400)
        self.assertEqual(self.speed(client, sid, 3, 0.69).status_code, 400)
        self.assertEqual(self.speed(client, sid, 3, "fast").status_code, 400)
        self.assertEqual(self.speed(client, sid, 99, 1.1).status_code, 404)

    # -- Block 6: export ------------------------------------------------------------------------
    def export(self, client, sid):
        response = client.get(f"/session/{sid}/export")
        self.assertEqual(response.status_code, 200, response.data[:300])
        self.assertIn("ep1_qc.zip", response.headers["Content-Disposition"])
        data = response.data
        response.close()
        return zipfile.ZipFile(io.BytesIO(data))

    def test_export_matches_ad_version_2_layout_and_writes_edits_into_the_excel(self):
        from openpyxl import load_workbook
        client, sid, _record = self.dubbed()
        self.regenerate(client, sid, 3, "[Angry] Get out!")
        self.speed(client, sid, 6, 1.1)
        archive = self.export(client, sid)
        names = sorted(archive.namelist())
        self.assertEqual(names, ["ep1_amrita.wav", "ep1_kuhu.wav", "ep1_timeline.csv",
                                 "processing_report.csv", "qc_changes.csv", "script_qc.xlsx"])
        # Tracks are exactly what the reviewer heard.
        self.assertEqual(archive.read("ep1_amrita.wav"), self.track_bytes(sid, 1))
        # V2's report columns, with the reviewed timings.
        import csv
        from v2.exporter import REPORT_FIELDS
        reader = csv.DictReader(io.StringIO(archive.read("processing_report.csv").decode("utf-8-sig")))
        report = {int(r["excel_row"]): r for r in reader}
        self.assertEqual(tuple(reader.fieldnames), REPORT_FIELDS)
        self.assertEqual(sorted(report), [2, 3, 4, 5, 6])
        self.assertEqual(report[2]["fit_speed"], "0.70")                          # pass 2's fit
        self.assertEqual(report[6]["fit_speed"], "1.10")                          # the reviewer's speed
        self.assertEqual(report[6]["pass1_duration_ms"], str(natural_ms(ONE_EPISODE[4][6])))  # the take
        # The Excel: edited text in TARGET DIALOGUE, highlighted with a note; other rows untouched.
        with archive.open("script_qc.xlsx") as handle:
            sheet = load_workbook(io.BytesIO(handle.read())).active
        self.assertEqual(sheet.cell(row=3, column=7).value, "[Angry] Get out!")
        self.assertIn("Script wording: [Angry] You beggar", sheet.cell(row=3, column=7).comment.text)
        self.assertIn("1.10x", sheet.cell(row=6, column=7).comment.text)
        self.assertIn("fitted speed was 0.81x", sheet.cell(row=6, column=7).comment.text)
        self.assertEqual(sheet.cell(row=2, column=7).value, ONE_EPISODE[0][6])
        self.assertIsNone(sheet.cell(row=2, column=7).comment)
        changes = archive.read("qc_changes.csv").decode("utf-8-sig").splitlines()
        self.assertEqual(len(changes), 3)  # header + row 3 (text) + row 6 (speed)
        self.assertIn("1.10x", changes[2])

    def test_export_without_changes_keeps_the_script_as_is(self):
        from openpyxl import load_workbook
        client, sid, _record = self.dubbed()
        archive = self.export(client, sid)
        with archive.open("script_qc.xlsx") as handle:
            sheet = load_workbook(io.BytesIO(handle.read())).active
        self.assertEqual([sheet.cell(row=r, column=7).value for r in range(2, 7)], [r[6] for r in ONE_EPISODE])
        self.assertEqual(len(archive.read("qc_changes.csv").decode("utf-8-sig").splitlines()), 1)

    def test_export_needs_a_finished_dub(self):
        gate = threading.Event()
        client = make_client(self.runtime, lambda t, v: (gate.wait(5), fake_voice(t, v))[1])
        sid = self.start(client).get_json()["session_id"]
        self.assertEqual(client.get(f"/session/{sid}/export").status_code, 409)
        gate.set()
        self.wait(client, sid)

    def test_bad_inputs(self):
        client = make_client(self.runtime, self.voice)
        sid = self.start(client).get_json()["session_id"]
        data = {"video": (io.BytesIO(b"x"), "episode.avi")}
        self.assertEqual(client.post(f"/session/{sid}/video", data=data,
                                     content_type="multipart/form-data").status_code, 400)
        self.assertEqual(client.get("/session/" + "0" * 32).status_code, 404)
        self.assertEqual(client.get("/session/../etc/video").status_code, 404)
        self.assertEqual(client.get(f"/session/{sid}/tracks/9.wav").status_code, 404)
        self.wait(client, sid)


@unittest.skipUnless(have_pydub(), "AD version 2 runs are built with pydub (Python 3.11/3.12)")
class ImportTests(unittest.TestCase):
    """Reviewing an AD version 2 run: its fitted line clips are the takes; nothing is dubbed again."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.runtime = Path(self.tmp.name) / "runtime"
        self.v2_runtime = Path(self.tmp.name) / "v2_runtime"
        self.calls = []
        self.client = make_client(self.runtime, self.voice, self.v2_runtime)

    def voice(self, text, voice_id):
        self.calls.append(text)
        return fake_voice(text, voice_id)

    def server_run(self, rows=ONE_EPISODE + SECOND_EPISODE, status="done", fps=25, finished=100.0):
        """A finished V2 job folder, as AD version 2 leaves it on the server."""
        run_id = os.urandom(16).hex()
        folder = self.v2_runtime / run_id
        folder.mkdir(parents=True)
        (folder / "AD_v2_output.zip").write_bytes(dub_zip(self.tmp.name, rows, name=f"{run_id}.zip", fps=fps))
        (folder / "run.json").write_text(json.dumps({
            "download_name": "Show_S1.zip", "created": finished - 10, "finished": finished, "fps": fps,
            "sheet": None, "rows": len(rows), "status": status, "episodes": sorted({r[0] for r in rows})}))
        return run_id

    def start(self, run_id, episode="1", **form):
        data = {"source": "v2", "run": run_id, "episode": episode, "fps": "25", **form}
        return self.client.post("/sessions", data=data, content_type="multipart/form-data")

    def wait(self, sid, timeout=30):
        deadline = time.time() + timeout
        while time.time() < deadline:
            record = self.client.get(f"/session/{sid}").get_json()
            if record["status"] != "dubbing":
                return record
            time.sleep(0.05)
        self.fail("import did not finish")

    def test_lists_finished_runs_newest_first(self):
        older = self.server_run(rows=ONE_EPISODE, finished=100.0)
        newer = self.server_run(finished=200.0)
        self.server_run(rows=ONE_EPISODE, status="running")             # not finished: not listed
        (self.v2_runtime / "not-a-run").mkdir()
        runs = self.client.get("/v2-runs").get_json()["runs"]
        self.assertEqual([r["id"] for r in runs], [newer, older])
        self.assertEqual(runs[0]["episodes"], ["1", "2"])
        details = self.client.get(f"/v2-runs/{newer}").get_json()
        self.assertEqual([e["episode"] for e in details["episodes"]], ["1", "2"])
        self.assertEqual(details["episodes"][0]["lines"], 5)
        self.assertEqual(details["fps"], 25)

    def test_episode_from_a_server_run_matches_ad_version_2_exactly(self):
        run_id = self.server_run()
        response = self.start(run_id)
        self.assertEqual(response.status_code, 201, response.get_json())
        record = self.wait(response.get_json()["session_id"])
        self.assertEqual(record["status"], "dubbed", record.get("error"))
        self.assertEqual(self.calls, [])                                   # no ElevenLabs request
        self.assertEqual((record["source"], record["episode"], record["summary"]["lines"]), ("v2", "1", 5))
        folder = self.runtime / record["session_id"]
        with zipfile.ZipFile(self.v2_runtime / run_id / "AD_v2_output.zip") as archive:
            # The tracks rebuilt from the line clips are V2's tracks, sample for sample.
            for track in record["audio"]["tracks"]:
                with wave.open(io.BytesIO(archive.read(f"pass2/tracks/ep1_{track['character'].lower()}.wav"))) as wav:
                    expected = wav.readframes(wav.getnframes())
                with wave.open(str(folder / "tracks" / f"{track['index']}.wav")) as wav:
                    self.assertEqual(wav.readframes(wav.getnframes()), expected, track["character"])
            self.assertEqual((folder / "clips" / "4.wav").read_bytes(), archive.read("pass2/lines/ep1_kuhu_row4.wav"))
            self.assertEqual((folder / "clips" / "4.take.wav").read_bytes(), archive.read("pass1/lines/ep1_kuhu_row4.wav"))
        lines = {line["row"]: line for line in record["audio"]["lines"]}
        self.assertEqual({row: line["fit_speed"] for row, line in lines.items()},
                         {2: 0.7, 3: 0.7, 4: 1.5, 5: 0.7, 6: 0.81})     # from V2's report
        self.assertAlmostEqual(lines[4]["needed_speed"], 2.16, delta=0.02)

    def test_imported_line_speed_is_measured_from_the_natural_take(self):
        record = self.wait(self.start(self.server_run()).get_json()["session_id"])
        sid = record["session_id"]
        folder = self.runtime / sid
        fitted = (folder / "clips" / "4.wav").read_bytes()
        natural = natural_ms(next(l["text"] for l in record["audio"]["lines"] if l["row"] == 4))
        body = self.client.post(f"/session/{sid}/lines/4/speed", json={"speed": 1.0}).get_json()
        line = next(l for l in body["audio"]["lines"] if l["row"] == 4)
        self.assertAlmostEqual(line["audio_end_ms"] - line["audio_start_ms"], natural, delta=40)
        self.client.post(f"/session/{sid}/lines/4/speed", json={"speed": 1.5})
        self.assertEqual((folder / "clips" / "4.wav").read_bytes(), fitted)  # back to V2's clip exactly

    def test_regenerating_an_imported_line_uses_elevenlabs_once(self):
        record = self.wait(self.start(self.server_run()).get_json()["session_id"])
        sid = record["session_id"]
        body = self.client.post(f"/session/{sid}/lines/3", json={"text": "Out!"}).get_json()
        self.assertEqual(self.calls, ["Out!"])
        line = next(l for l in body["audio"]["lines"] if l["row"] == 3)
        self.assertEqual((line["fit_speed"], line["edited"]), (0.7, True))
        archive = zipfile.ZipFile(io.BytesIO(self.client.get(f"/session/{sid}/export").data))
        self.assertIn("ep1_kuhu.wav", archive.namelist())

    def test_uploaded_zip_finds_the_frame_rate_the_run_used(self):
        data = dub_zip(self.tmp.name, ONE_EPISODE + SECOND_EPISODE, name="thirty.zip", fps=30)
        form = {"dub": (io.BytesIO(data), "Show_S1.zip"), "fps": "25"}
        response = self.client.post("/v2-runs", data=form, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 201, response.get_json())
        details = response.get_json()
        self.assertEqual((details["fps"], details["name"]), (30, "Show_S1.zip"))
        self.assertTrue(details["id"].startswith("up-"))
        response = self.start(details["id"], episode="2")
        self.assertEqual(response.status_code, 201, response.get_json())
        record = self.wait(response.get_json()["session_id"])
        self.assertEqual((record["status"], record["fps"], record["summary"]["lines"]), ("dubbed", 30, 2))
        self.assertEqual([t["character"] for t in record["audio"]["tracks"]], ["KUHU", "RAVI"])

    def test_bad_runs_are_refused_with_a_reason(self):
        upload = lambda data, name="run.zip": self.client.post(
            "/v2-runs", data={"dub": (io.BytesIO(data), name), "fps": "25"}, content_type="multipart/form-data")
        self.assertIn("not a ZIP", upload(b"not a zip").get_json()["error"])
        self.assertEqual(upload(b"x", "run.xlsx").status_code, 400)
        old = io.BytesIO()
        with zipfile.ZipFile(old, "w") as archive:          # an earlier V2 ZIP: lines/ at the top, no passes
            archive.writestr("lines/1_KUHU_2.wav", b"RIFF")
        response = upload(old.getvalue())
        self.assertEqual(response.status_code, 400)
        self.assertIn("earlier AD Tool V2", response.get_json()["error"])
        self.assertFalse(any((self.runtime / "uploads").iterdir()))      # refused uploads are not kept

        run_id = self.server_run(rows=ONE_EPISODE)
        self.assertEqual(self.start(run_id, episode="7").status_code, 400)
        self.assertEqual(self.start("0" * 32).status_code, 404)
        self.assertEqual(self.client.get("/v2-runs/../../etc").status_code, 404)
        # A clip missing from the ZIP is reported by row before anything is stored.
        path = self.v2_runtime / run_id / "AD_v2_output.zip"
        with zipfile.ZipFile(path) as archive:
            items = {name: archive.read(name) for name in archive.namelist() if name != "pass2/lines/ep1_amrita_row3.wav"}
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in items.items():
                archive.writestr(name, data)
        response = self.start(run_id)
        self.assertEqual(response.status_code, 400)
        self.assertIn("Excel row 3", response.get_json()["error"])
        self.assertFalse(any(p.name != "uploads" for p in self.runtime.iterdir()))


if __name__ == "__main__":
    unittest.main()
